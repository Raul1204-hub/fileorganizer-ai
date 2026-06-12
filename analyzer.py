import base64
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Callable

import openpyxl
import pdfplumber

from config import ANALYSIS_MODEL, OCR_LANGUAGES, OCR_MIN_CHARS, VISION_MAX_MB, VISION_MAX_SIDE, VISION_MODEL
from log import get_logger
from ollama_client import call_ollama, call_ollama_vision
from scanner import IMG_EXTS

try:
    from PIL import Image as _PIL_Image

    _PILLOW_AVAILABLE = True
except ImportError:
    _PIL_Image = None  # type: ignore[assignment]
    _PILLOW_AVAILABLE = False

_TESSERACT_AVAILABLE: bool | None = None

logger = get_logger("fileorganizer.analyzer")

_DOCX_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# ── Validation constants ──────────────────────────────────────────────────────

_VALID_CATEGORIAS = frozenset(
    {
        "Documentos",
        "Imágenes",
        "Audio",
        "Vídeo",
        "Código",
        "Datos",
        "Comprimidos",
        "Programas",
        "Desconocido",
    }
)

_ANALYSIS_SCHEMA = {
    "type": "object",
    "properties": {
        "categoria": {"type": "string"},
        "etiquetas": {"type": "array", "items": {"type": "string"}},
        "resumen": {"type": "string"},
    },
    "required": ["categoria", "etiquetas", "resumen"],
}

# .doc legacy heuristic: extracted text must have this fraction of alphabetic chars
_DOC_ALPHA_RATIO_MIN = 0.45

_VISION_PROMPT = (
    "Analyze this image. Respond with ONLY a JSON object:\n"
    '{"categoria": "...", "etiquetas": ["tag1", ...], "resumen": "..."}\n\n'
    "categoria options: Imágenes, Documentos, Código, Datos, Desconocido\n"
    "etiquetas: up to 5 descriptive tags in Spanish\n"
    "resumen: 1-2 sentences in Spanish describing what the image shows\n"
    "Output JSON only, no explanation."
)


# ── Vision helpers ────────────────────────────────────────────────────────────


def _resize_image(path: Path, max_side: int) -> bytes | None:
    """Load image, resize so the longest side <= max_side, return JPEG bytes."""
    if not _PILLOW_AVAILABLE:
        logger.warning("vision | Pillow not installed — pip install Pillow")
        return None
    import io

    try:
        with _PIL_Image.open(path) as img:
            img = img.convert("RGB")
            w, h = img.size
            if max(w, h) > max_side:
                ratio = max_side / max(w, h)
                img = img.resize((int(w * ratio), int(h * ratio)), _PIL_Image.LANCZOS)
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=85)
            return buf.getvalue()
    except Exception as exc:
        logger.warning("image_resize | %s | %s", path.name, exc)
        return None


def analyze_image(
    path: Path,
    on_failure: Callable[[str], None] | None = None,
) -> dict:
    """Analyze an image with the configured vision model via Ollama.

    Returns {categoria, etiquetas, resumen, _text_len=0, _texto_via='vision'}
    or {} on failure (model missing, too large, resize error).
    """
    size_mb = path.stat().st_size / (1024 * 1024)
    if size_mb > VISION_MAX_MB:
        logger.debug("vision | skip_large | %s | %.1f MB > %.0f MB", path.name, size_mb, VISION_MAX_MB)
        if on_failure:
            on_failure("extraction")
        return {}

    img_bytes = _resize_image(path, VISION_MAX_SIDE)
    if not img_bytes:
        if on_failure:
            on_failure("extraction")
        return {}

    b64 = base64.b64encode(img_bytes).decode("ascii")

    for _attempt in range(2):
        raw = call_ollama_vision(VISION_MODEL, _VISION_PROMPT, b64, fmt="json")
        if not raw:
            continue
        try:
            data = json.loads(raw)
        except json.JSONDecodeError as exc:
            logger.warning("vision_parse | %s | attempt=%d | %s", path.name, _attempt + 1, exc)
            continue
        result = _validate_analysis(data)
        if result and result.get("resumen"):
            result["_texto_via"] = "vision"
            result["_text_len"] = 0
            return result

    logger.warning("vision_failed | %s", path.name)
    if on_failure:
        on_failure("ollama")
    return {}


# ── OCR helpers ───────────────────────────────────────────────────────────────


def _check_tesseract() -> bool:
    """Check tesseract availability once and cache the result."""
    global _TESSERACT_AVAILABLE
    if _TESSERACT_AVAILABLE is not None:
        return _TESSERACT_AVAILABLE
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
        _TESSERACT_AVAILABLE = True
    except Exception:
        logger.warning(
            "ocr | Tesseract not found in PATH — OCR disabled. "
            "Install from: https://github.com/UB-Mannheim/tesseract/wiki"
        )
        _TESSERACT_AVAILABLE = False
    return _TESSERACT_AVAILABLE


def _ocr_pdf(path: Path) -> str:
    """Render the first 3 pages of a PDF with pypdfium2 and OCR with pytesseract.

    Returns '' if any dependency is missing or rendering fails.
    """
    try:
        import pypdfium2 as pdfium
    except ImportError:
        logger.warning("ocr | pypdfium2 not installed — pip install pypdfium2")
        return ""
    try:
        import pytesseract
    except ImportError:
        return ""

    texts: list[str] = []
    try:
        doc = pdfium.PdfDocument(str(path))
        pages_to_ocr = min(len(doc), 3)
        if pages_to_ocr == 0:
            doc.close()
            return ""
        for i in range(pages_to_ocr):
            page = doc[i]
            bitmap = page.render(scale=200 / 72)
            pil_img = bitmap.to_pil()
            texts.append(pytesseract.image_to_string(pil_img, lang=OCR_LANGUAGES))
        doc.close()
    except Exception as exc:
        logger.warning("ocr | pdf_render | %s | %s", path.name, exc)
        return ""

    return "\n\n".join(texts).strip()[:4000]


def _try_ocr_pdf(path: Path) -> str:
    """Run OCR only when tesseract is available; return '' otherwise."""
    if not _check_tesseract():
        return ""
    return _ocr_pdf(path)


# ── text extraction ───────────────────────────────────────────────────────────


def extract_text_pdf(path: Path) -> str:
    parts: list[str] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages[:10]:
                text = page.extract_text()
                if text:
                    parts.append(text)
    except Exception as e:
        logger.warning("extract_pdf | %s | %s", path, e)
    return "\n".join(parts)[:4000]


def extract_text_docx(path: Path) -> str:
    """Extract text from modern OOXML .docx (ZIP-based)."""
    try:
        with zipfile.ZipFile(path) as z:
            if "word/document.xml" not in z.namelist():
                return ""
            with z.open("word/document.xml") as f:
                root = ET.parse(f).getroot()
                texts = [node.text for node in root.iter(f"{{{_DOCX_NS}}}t") if node.text]
                return " ".join(texts)[:4000]
    except Exception as e:
        logger.warning("extract_docx | %s | %s", path, e)
        return ""


def extract_text_doc_legacy(path: Path) -> str:
    """Best-effort extraction for legacy OLE2 .doc binary format.

    OLE2 is not ZIP-based so zipfile cannot read it. Instead, decode the file
    as latin-1 and extract printable sequences (like the UNIX 'strings' tool).
    Apply an alphabetic-ratio heuristic to discard noise.
    Always logs a warning because extraction quality is degraded.
    """
    try:
        raw = path.read_bytes()
        # latin-1 decode never fails: every byte maps to a code point
        text = raw.decode("latin-1")
        # Printable ASCII + Latin-1 supplement, minimum 5 chars
        sequences = re.findall(r"[ -~\xa0-\xff]{5,}", text)
        result = " ".join(s.strip() for s in sequences if s.strip())
    except Exception as e:
        logger.warning("extract_doc_legacy | %s | %s", path, e)
        return ""

    if not result:
        return ""

    alpha_count = sum(1 for c in result if c.isalpha())
    ratio = alpha_count / len(result)

    if ratio < _DOC_ALPHA_RATIO_MIN:
        logger.warning(
            "extract_doc_legacy | %s | discarded — alpha ratio %.2f below threshold (likely binary noise)",
            path,
            ratio,
        )
        return ""

    logger.warning(
        "extract_doc_legacy | %s | degraded OLE2 extraction — alpha %.0f%%, %d chars extracted",
        path,
        ratio * 100,
        len(result),
    )
    return result[:4000]


def extract_text_xlsx(path: Path) -> str:
    """Extract cell values from .xlsx using openpyxl.

    Reads the first 100 rows of the active sheet in read-only / data-only mode
    so formula results (numbers, dates) are included, not the formula strings.
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            ws = wb.worksheets[0] if wb.worksheets else None
        if ws is None:
            wb.close()
            return ""

        rows: list[str] = []
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i > 100:
                break
            cells = [str(c) if c is not None else "" for c in row]
            line = "\t".join(cells).strip()
            if line:
                rows.append(line)

        wb.close()
        return "\n".join(rows)[:4000]
    except Exception as e:
        logger.warning("extract_xlsx | %s | %s", path, e)
        return ""


def extract_text_plain(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="ignore")[:4000]
    except Exception as e:
        logger.warning("extract_plain | %s | %s", path, e)
        return ""


def extract_text(path: Path, extension: str) -> str:
    ext = extension.lower()
    if ext == ".pdf":
        return extract_text_pdf(path)
    if ext == ".docx":
        return extract_text_docx(path)
    if ext == ".doc":
        # Legacy OLE2 binary format — separate degraded extractor
        return extract_text_doc_legacy(path)
    if ext == ".xlsx":
        return extract_text_xlsx(path)
    if ext in (".txt", ".odt", ".csv"):
        return extract_text_plain(path)
    return ""


# ── Validation ────────────────────────────────────────────────────────────────


def _validate_analysis(data) -> dict:
    """Normalize and validate an Ollama analysis result.

    Returns a 3-key dict on success, or {} if data is not a dict.
    """
    if not isinstance(data, dict):
        return {}

    categoria = data.get("categoria", "")
    if categoria not in _VALID_CATEGORIAS:
        categoria = "Desconocido"

    raw_tags = data.get("etiquetas", [])
    if isinstance(raw_tags, str):
        raw_tags = [raw_tags]
    elif not isinstance(raw_tags, list):
        raw_tags = []
    etiquetas = [e.strip() for e in raw_tags if isinstance(e, str) and e.strip()][:5]

    resumen = data.get("resumen", "")
    if not isinstance(resumen, str):
        resumen = str(resumen)
    resumen = resumen.strip()[:300]

    return {"categoria": categoria, "etiquetas": etiquetas, "resumen": resumen}


# ── Ollama analysis ───────────────────────────────────────────────────────────


def analyze_with_ollama(text: str, filename: str) -> dict:
    """Call Ollama with JSON Schema format and validate.

    Retries once on parse/validation failure.
    Returns {} if both attempts fail or Ollama is unavailable.
    """
    prompt = (
        f'Analyze this document named "{filename}". '
        "Respond in JSON with keys: categoria, etiquetas (list of tags), resumen.\n\n"
        f"Document content:\n{text[:2500]}"
    )

    for _attempt in range(2):
        try:
            raw = call_ollama(ANALYSIS_MODEL, prompt, fmt=_ANALYSIS_SCHEMA)
        except RuntimeError as e:
            logger.warning("ollama_call | %s | %s", filename, e)
            return {}
        if not raw:
            continue

        try:
            data = json.loads(raw)
        except json.JSONDecodeError as e:
            logger.warning("ollama_parse | %s | attempt=%d | %s", filename, _attempt + 1, e)
            continue

        result = _validate_analysis(data)
        if result and len(result) == 3:
            return result

    logger.warning("ollama_failed | %s | both attempts exhausted", filename)
    return {}


# ── public API ────────────────────────────────────────────────────────────────


def analyze_file(
    path: Path,
    extension: str,
    on_failure: Callable[[str], None] | None = None,
) -> dict:
    """Extract text (or use vision) and analyze with Ollama.

    Parameters
    ----------
    on_failure : optional callback(reason: str) where reason is
        'extraction' (no/empty text, image too large, resize failed) or
        'ollama' (model returned nothing useful).

    Returns a dict with {categoria, etiquetas, resumen, _text_len, _texto_via}
    on success, or {} on failure.  _texto_via is 'pdf', 'ocr', or 'vision';
    absent for non-PDF text documents.
    """
    ext = extension.lower()

    # ── Image path (vision model) ─────────────────────────────────────────────
    if ext in IMG_EXTS:
        return analyze_image(path, on_failure=on_failure)

    # ── Document path ─────────────────────────────────────────────────────────
    text = extract_text(path, extension)
    texto_via: str | None = None

    if ext == ".pdf":
        texto_via = "pdf"
        if len(text.strip()) < OCR_MIN_CHARS:
            ocr_text = _try_ocr_pdf(path)
            if ocr_text:
                text = ocr_text
                texto_via = "ocr"

    if not text.strip():
        if on_failure:
            on_failure("extraction")
        return {}

    result = analyze_with_ollama(text, path.name)
    if not result:
        if on_failure:
            on_failure("ollama")
        return {}

    result["_text_len"] = len(text)
    if texto_via is not None:
        result["_texto_via"] = texto_via
    return result
